import csv
import io
import re

SUPPORTED_EXTENSIONS = {"pdf", "docx", "txt", "md", "csv", "xlsx"}


def clean_text(value):
    value = str(value or "").replace("\x00", " ")
    value = re.sub(r"[ \t]+", " ", value)
    return re.sub(r"\n{3,}", "\n\n", value).strip()


def chunk_text(text, max_chars=3500):
    paragraphs = [clean_text(item) for item in re.split(r"\n\s*\n", text) if clean_text(item)]
    chunks, current = [], []
    for paragraph in paragraphs:
        if current and sum(len(item) for item in current) + len(paragraph) > max_chars:
            chunks.append("\n\n".join(current)); current = []
        if len(paragraph) > max_chars:
            for start in range(0, len(paragraph), max_chars): chunks.append(paragraph[start:start + max_chars])
        else: current.append(paragraph)
    if current: chunks.append("\n\n".join(current))
    return chunks


def parse_document(data, extension):
    extension = extension.lower().lstrip(".")
    if extension not in SUPPORTED_EXTENSIONS: raise ValueError("Unsupported document type")
    passages, page_count = [], None
    if extension == "pdf":
        from pypdf import PdfReader
        reader = PdfReader(io.BytesIO(data)); page_count = len(reader.pages)
        for page_number, page in enumerate(reader.pages, 1):
            text = clean_text(page.extract_text() or "")
            for index, chunk in enumerate(chunk_text(text), 1):
                passages.append({"locator_type":"page","locator":f"page {page_number}, passage {index}","page_from":page_number,"page_to":page_number,"content":chunk})
    elif extension == "docx":
        from docx import Document
        document = Document(io.BytesIO(data)); paragraph_number = 0
        for paragraph in document.paragraphs:
            text = clean_text(paragraph.text)
            if text:
                paragraph_number += 1; passages.append({"locator_type":"paragraph","locator":f"paragraph {paragraph_number}","content":text})
        for table_number, table in enumerate(document.tables, 1):
            for row_number, row in enumerate(table.rows, 1):
                text = clean_text(" | ".join(cell.text for cell in row.cells))
                if text: passages.append({"locator_type":"table_row","locator":f"table {table_number}, row {row_number}","content":text})
    elif extension == "xlsx":
        from openpyxl import load_workbook
        workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        for sheet in workbook.worksheets:
            for row_number, row in enumerate(sheet.iter_rows(values_only=True), 1):
                text = clean_text(" | ".join("" if value is None else str(value) for value in row))
                if text: passages.append({"locator_type":"row","locator":f"sheet {sheet.title}, row {row_number}","content":text})
    elif extension == "csv":
        decoded = data.decode("utf-8-sig", errors="replace")
        for row_number, row in enumerate(csv.reader(io.StringIO(decoded)), 1):
            text = clean_text(" | ".join(row))
            if text: passages.append({"locator_type":"row","locator":f"row {row_number}","content":text})
    else:
        decoded = data.decode("utf-8-sig", errors="replace")
        for index, chunk in enumerate(chunk_text(clean_text(decoded)), 1): passages.append({"locator_type":"passage","locator":f"passage {index}","content":chunk})
    if not passages: raise ValueError("No extractable text was found in this document")
    if len(passages) > 5000: raise ValueError("Document contains too many passages")
    return {"passages": passages, "page_count": page_count, "passage_count": len(passages)}


if __name__ == "__main__":
    import base64, json, sys
    payload=json.load(sys.stdin)
    print(json.dumps(parse_document(base64.b64decode(payload["data"],validate=True),payload["extension"])))
