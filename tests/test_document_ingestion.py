import io
import pathlib
import sys
import unittest

ROOT=pathlib.Path(__file__).resolve().parents[1]
sys.path.insert(0,str(ROOT / "scripts"))
from parse_legal_document import parse_document


class DocumentIngestionTests(unittest.TestCase):
    def test_pdf_preserves_page_anchors(self):
        from reportlab.pdfgen.canvas import Canvas
        output=io.BytesIO();canvas=Canvas(output);canvas.drawString(70,750,"Claim one requires a processor.");canvas.showPage();canvas.drawString(70,750,"The cited reference lacks that processor.");canvas.save()
        parsed=parse_document(output.getvalue(),"pdf")
        self.assertEqual(parsed["page_count"],2)
        self.assertEqual(parsed["passages"][0]["page_from"],1)
        self.assertEqual(parsed["passages"][1]["page_from"],2)

    def test_docx_preserves_paragraph_and_table_anchors(self):
        from docx import Document
        document=Document();document.add_paragraph("Governing law is England and Wales.");table=document.add_table(rows=1,cols=2);table.cell(0,0).text="Claim element";table.cell(0,1).text="Evidence";output=io.BytesIO();document.save(output)
        parsed=parse_document(output.getvalue(),"docx")
        self.assertEqual(parsed["passages"][0]["locator"],"paragraph 1")
        self.assertEqual(parsed["passages"][1]["locator"],"table 1, row 1")

    def test_xlsx_and_csv_have_stable_row_anchors(self):
        from openpyxl import Workbook
        workbook=Workbook();workbook.active.append(["Patent","Status"]);workbook.active.append(["EP123","Live"]);output=io.BytesIO();workbook.save(output)
        xlsx=parse_document(output.getvalue(),"xlsx");csv=parse_document(b"Mark,Class\nSALLY,9\n","csv")
        self.assertEqual(xlsx["passages"][1]["locator"],"sheet Sheet, row 2")
        self.assertEqual(csv["passages"][1]["locator"],"row 2")

    def test_empty_document_is_rejected(self):
        with self.assertRaisesRegex(ValueError,"No extractable text"):
            parse_document(b"   ","txt")


if __name__=="__main__":unittest.main()
