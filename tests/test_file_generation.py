import io
import pathlib
import sys
import unittest

ROOT = pathlib.Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "scripts"))

from create_chat_file import build_file


class FileGenerationTests(unittest.TestCase):
    def test_pdf_contains_document_not_capability_disclaimer(self):
        from pypdf import PdfReader
        _, mime, data = build_file("pdf", "IP Licence Agreement", "# Agreement\nThe Licensor grants the Licence.")
        text = "\n".join(page.extract_text() or "" for page in PdfReader(io.BytesIO(data)).pages).lower()
        self.assertEqual(mime, "application/pdf")
        self.assertIn("the licensor grants", text)
        self.assertNotIn("i can't generate", text)
        self.assertNotIn("google docs", text)

    def test_capability_statement_is_rejected(self):
        with self.assertRaisesRegex(ValueError, "capability"):
            build_file("pdf", "Bad", "I cannot create a PDF. Copy this into Word.")

    def test_docx_has_native_heading_and_paragraph(self):
        from docx import Document
        _, mime, data = build_file("docx", "Trademark Report", "# Findings\nNo conflicting mark was identified.")
        doc = Document(io.BytesIO(data))
        self.assertIn("Findings", [paragraph.text for paragraph in doc.paragraphs])
        self.assertIn("No conflicting mark was identified.", [paragraph.text for paragraph in doc.paragraphs])
        self.assertIn("wordprocessingml", mime)

    def test_safe_filename_blocks_paths_and_duplicate_extensions(self):
        name, _, _ = build_file("txt", "../../Agreement.pdf.exe", "Substantive content")
        self.assertEqual(name, "agreement-pdf-exe.txt")
        self.assertNotIn("..", name)
        self.assertNotIn("/", name)

    def test_text_exports_preserve_substantive_content(self):
        content = "# Governing Law\nThis Agreement is governed by the laws of England and Wales."
        for fmt in ("txt", "md"):
            _, _, data = build_file(fmt, "Agreement", content)
            self.assertEqual(data.decode("utf-8"), content)

    def test_same_artifact_exports_deterministically(self):
        content = "# Licence\nThe licence is non-exclusive and non-transferable."
        first = build_file("pdf", "IP Licence", content)[2]
        second = build_file("pdf", "IP Licence", content)[2]
        from pypdf import PdfReader
        def text(data):
            return "\n".join(page.extract_text() or "" for page in PdfReader(io.BytesIO(data)).pages)
        self.assertEqual(text(first), text(second))

    def test_requested_filename_is_sanitized_and_format_controlled(self):
        name, _, _ = build_file("pdf", "Agreement", "Substantive terms", "../Client Final.exe.pdf")
        self.assertEqual(name, "client-final-exe-pdf.pdf")

    def test_csv_and_xlsx_are_real_structured_files(self):
        from openpyxl import load_workbook
        _, csv_mime, csv_data = build_file("csv", "Claim Chart", "# Element 1\nA processor")
        self.assertIn("text/csv", csv_mime)
        self.assertIn("Element 1", csv_data.decode("utf-8-sig"))
        _, xlsx_mime, xlsx_data = build_file("xlsx", "Claim Chart", "# Element 1\nA processor")
        workbook = load_workbook(io.BytesIO(xlsx_data))
        self.assertEqual(workbook.active["A1"].value, "Claim Chart")
        self.assertIn("spreadsheetml", xlsx_mime)


    def test_markdown_bold_and_dividers_render_cleanly(self):
        from pypdf import PdfReader
        from docx import Document

        raw_md = """# EUROPEAN PATENT APPLICATION
**Statutory Authority**: EPC Article 75
**Target Jurisdiction**: EPO
---
## 1. TITLE
**AUTOMATED DRONE BATTERY SWAPPING**
"""
        # Test PDF
        _, _, pdf_data = build_file("pdf", "European Patent", raw_md)
        pdf_text = "\n".join(page.extract_text() or "" for page in PdfReader(io.BytesIO(pdf_data)).pages)
        self.assertNotIn("**", pdf_text, "PDF must not contain raw markdown asterisks '**'")
        self.assertIn("Statutory Authority", pdf_text)
        self.assertIn("AUTOMATED DRONE BATTERY SWAPPING", pdf_text)

        # Test DOCX
        _, _, docx_data = build_file("docx", "European Patent", raw_md)
        doc = Document(io.BytesIO(docx_data))
        all_para_text = [p.text for p in doc.paragraphs]
        for t in all_para_text:
            self.assertNotIn("**", t, f"DOCX paragraph must not contain raw asterisks: {t}")
        # Check that bold runs exist
        all_runs = [r for p in doc.paragraphs for r in p.runs]
        bold_runs = [r.text for r in all_runs if r.bold]
        self.assertTrue(any("Statutory Authority" in b for b in bold_runs), "Expected 'Statutory Authority' to be bold run in DOCX")


if __name__ == "__main__":
    unittest.main()
